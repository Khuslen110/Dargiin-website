import csv
import io

from openpyxl import load_workbook

from tests.conftest import auth_headers, create_person, login
from tests.test_people import _create_scoped_user, _login_scoped

COLUMNS = [
    "Регистрийн дугаар",
    "Овог",
    "Нэр",
    "Хүйс",
    "Төрсөн огноо",
    "Утас",
    "Дүүрэг",
    "Хороо",
    "Хаяг",
    "Ажил",
    "Бүлгийн төрөл",
    "Харьяалагдах хүний регистр",
]


def _csv(rows: list[list[str]], columns: list[str] | None = None) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(columns or COLUMNS)
    writer.writerows(rows)
    return buffer.getvalue()


def _upload(
    client,
    csrf,
    content: str,
    filename: str = "import.csv",
    parent_id: int | None = None,
):
    data = {"parent_id": str(parent_id)} if parent_id is not None else None
    return client.post(
        "/api/imports/preview",
        files={"file": (filename, content.encode("utf-8"), "text/csv")},
        data=data,
        headers=auth_headers(csrf),
    )


def test_preview_confirm_and_hierarchy(admin_auth):
    client, csrf = admin_auth["client"], admin_auth["csrf"]
    content = _csv(
        [
            [
                "ТБ10000001",
                "Бат",
                "Сүх",
                "Эрэгтэй",
                "1990-01-01",
                "99110001",
                "Хан-Уул",
                "1-р хороо",
                "Хаяг",
                "Инженер",
                "Түмт",
                "",
            ],
            [
                "ТБ10000002",
                "Бат",
                "Дорж",
                "Эрэгтэй",
                "1995-02-02",
                "99110002",
                "Хан-Уул",
                "2-р хороо",
                "Хаяг",
                "Багш",
                "Мянгат",
                "ТБ10000001",
            ],
            ["ТБ10000002", "Давхар", "Мөр", "Эмэгтэй", "1990-01-01", "", "", "", "", "", "Зуут", ""],
            ["БУРУУ", "Алдаа", "Мөр", "Эрэгтэй", "1990-01-01", "", "", "", "", "", "Зуут", ""],
        ]
    )
    preview = _upload(client, csrf, content)
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["total_rows"] == 4
    assert body["valid_rows"] == 2
    assert body["invalid_rows"] == 2
    assert body["can_confirm"] is True
    assert len(body["errors"]) == 2
    job_id = body["job_id"]

    confirm = client.post(f"/api/imports/{job_id}/confirm", headers=auth_headers(csrf))
    assert confirm.status_code == 200, confirm.text
    result = confirm.json()
    assert result["status"] == "COMPLETED"
    assert result["imported_rows"] == 2

    child = client.get("/api/people", params={"reg_number": "ТБ10000002"}).json()["items"][0]
    detail = client.get(f"/api/people/{child['id']}").json()
    assert detail["path"][0]["reg_number"] == "ТБ10000001"

    jobs = client.get("/api/imports").json()
    assert jobs["total"] == 1


def test_import_under_selected_parent(admin_auth):
    client, csrf = admin_auth["client"], admin_auth["csrf"]
    root = create_person(client, csrf)
    columns = [column for column in COLUMNS if column != "Харьяалагдах хүний регистр"]
    content = _csv(
        [
            [
                "ТБ30000001",
                "Одоо",
                "Нэг",
                "Эрэгтэй",
                "1990-01-01",
                "",
                "",
                "",
                "",
                "",
                "Мянгат",
            ],
            [
                "ТБ30000002",
                "Одоо",
                "Хоёр",
                "Эмэгтэй",
                "1991-02-02",
                "",
                "",
                "",
                "",
                "",
                "Мянгат",
            ],
        ],
        columns,
    )
    preview = _upload(client, csrf, content, parent_id=root["id"])
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["valid_rows"] == 2
    assert body["invalid_rows"] == 0

    confirm = client.post(
        f"/api/imports/{body['job_id']}/confirm", headers=auth_headers(csrf)
    )
    assert confirm.status_code == 200, confirm.text
    assert confirm.json()["imported_rows"] == 2

    for reg in ("ТБ30000001", "ТБ30000002"):
        item = client.get("/api/people", params={"reg_number": reg}).json()["items"][0]
        detail = client.get(f"/api/people/{item['id']}").json()
        assert detail["path"][0]["reg_number"] == root["reg_number"]


def test_import_header_validation(admin_auth):
    client, csrf = admin_auth["client"], admin_auth["csrf"]
    preview = _upload(client, csrf, "Нэр,Овог\nБат,Дорж\n")
    assert preview.status_code == 400
    assert "толгой" in preview.json()["detail"].lower()


def test_import_is_admin_only(admin_auth):
    client = admin_auth["client"]
    root = create_person(client, admin_auth["csrf"])
    _create_scoped_user(admin_auth, root["id"], username="importer")
    csrf = _login_scoped(client, "importer")
    content = _csv(
        [
            [
                "ТБ20000001",
                "Хүү",
                "Мөр",
                "Эрэгтэй",
                "1990-01-01",
                "",
                "",
                "",
                "",
                "",
                "Мянгат",
                root["reg_number"],
            ],
        ]
    )
    assert _upload(client, csrf, content).status_code == 403


def test_error_report_and_template(admin_auth):
    client, csrf = admin_auth["client"], admin_auth["csrf"]
    preview = _upload(
        client,
        csrf,
        _csv(
            [
                ["БУРУУ", "Алдаа", "Мөр", "Эрэгтэй", "", "", "", "", "", "", "Зуут", ""],
            ]
        ),
    )
    assert preview.status_code == 200
    job_id = preview.json()["job_id"]
    errors = client.get(f"/api/imports/{job_id}/errors.csv")
    assert errors.status_code == 200
    assert "Мөрийн дугаар" in errors.content.decode("utf-8-sig")
    template = client.get("/api/imports/template")
    assert template.status_code == 200
    template_text = template.content.decode("utf-8-sig")
    assert "Регистрийн дугаар" in template_text
    assert "Хэрэглэгчийн нэр" in template_text
    assert "Нууц үг" in template_text
    assert "Харьяалагдах хүний регистр" not in template_text


def test_export_csv_formula_injection_and_xlsx(admin_auth):
    client, csrf = admin_auth["client"], admin_auth["csrf"]
    create_person(
        client,
        csrf,
        last_name='=HYPERLINK("http://evil")',
        job="+SUM(1,1)",
    )
    csv_response = client.get("/api/exports/people", params={"fmt": "csv"})
    assert csv_response.status_code == 200
    text = csv_response.content.decode("utf-8-sig")
    assert "'=HYPERLINK" in text
    assert "'+SUM" in text
    assert csv_response.headers["X-Export-Count"] == "1"

    xlsx_response = client.get("/api/exports/people", params={"fmt": "xlsx"})
    assert xlsx_response.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx_response.content), read_only=True)
    sheet = workbook.active
    values = [cell for row in sheet.iter_rows(values_only=True) for cell in row]
    assert any(str(value).startswith("'=HYPERLINK") for value in values if value)
    workbook.close()


def test_export_is_admin_only(admin_auth):
    client = admin_auth["client"]
    root = create_person(client, admin_auth["csrf"])
    _create_scoped_user(admin_auth, root["id"], username="noexport")
    _login_scoped(client, "noexport")
    assert client.get("/api/exports/people").status_code == 403


ACCOUNT_COLUMNS = COLUMNS[:-1] + ["Хэрэглэгчийн нэр", "Нууц үг", COLUMNS[-1]]


def test_import_creates_accounts_from_password_column(admin_auth):
    client, csrf = admin_auth["client"], admin_auth["csrf"]
    root = create_person(client, csrf)
    content = _csv(
        [
            [
                "ТБ40000001",
                "Импорт",
                "Мянгат",
                "Эрэгтэй",
                "1990-01-01",
                "99114001",
                "Хан-Уул",
                "1-р хороо",
                "Хаяг",
                "Ажил",
                "Мянгат",
                "import.mgr",
                "ImportPass2026",
                root["reg_number"],
            ],
            [
                "ТБ40000002",
                "Импорт",
                "Зуут",
                "Эмэгтэй",
                "1991-02-02",
                "99114002",
                "Хан-Уул",
                "1-р хороо",
                "Хаяг",
                "Ажил",
                "Зуут",
                "import.zuut",
                "ImportPass2026",
                "ТБ40000001",
            ],
            [
                "ТБ40000003",
                "Импорт",
                "Аравт",
                "Эрэгтэй",
                "1992-03-03",
                "99114003",
                "Хан-Уул",
                "1-р хороо",
                "Хаяг",
                "Ажил",
                "Аравт",
                "import.aravt",
                "ImportPass2026",
                "ТБ40000002",
            ],
            [
                "ТБ40000004",
                "Авто",
                "Гишүүн",
                "Эмэгтэй",
                "1993-04-04",
                "99114004",
                "Хан-Уул",
                "1-р хороо",
                "Хаяг",
                "Ажил",
                "Гишүүн",
                "",
                "ImportPass2026",
                "ТБ40000003",
            ],
        ],
        ACCOUNT_COLUMNS,
    )
    preview = _upload(client, csrf, content)
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["valid_rows"] == 4
    assert body["accounts_planned"] == 4

    confirm = client.post(
        f"/api/imports/{body['job_id']}/confirm", headers=auth_headers(csrf)
    )
    assert confirm.status_code == 200, confirm.text
    result = confirm.json()
    assert result["imported_rows"] == 4
    assert result["accounts_created"] == 4

    member = client.get("/api/people", params={"reg_number": "ТБ40000004"}).json()["items"][0]
    listing = client.get("/api/users", params={"search": "Авто"}).json()
    assert listing["total"] == 1
    auto_user = listing["items"][0]
    assert auto_user["username"] == f"user{member['id']}"
    assert auto_user["role"] == "MEMBER"

    from fastapi.testclient import TestClient

    from app.main import app

    imported_client = TestClient(app)
    login(imported_client, "import.mgr", "ImportPass2026")
    me = imported_client.get("/api/auth/me").json()["user"]
    assert me["role"] == "LEADER"
    assert me["leader_person"]["id"] == client.get(
        "/api/people", params={"reg_number": "ТБ40000001"}
    ).json()["items"][0]["id"]

    member_client = TestClient(app)
    login(member_client, f"user{member['id']}", "ImportPass2026")
    assert member_client.get("/api/auth/me").json()["user"]["role"] == "MEMBER"


def test_import_password_validation_does_not_echo_password(admin_auth):
    client, csrf = admin_auth["client"], admin_auth["csrf"]
    content = _csv(
        [
            [
                "ТБ50000001",
                "Алдаа",
                "Нэг",
                "Эрэгтэй",
                "",
                "",
                "",
                "",
                "",
                "",
                "Түмт",
                "bad.user",
                "short",
                "",
            ],
            [
                "ТБ50000002",
                "Алдаа",
                "Хоёр",
                "Эрэгтэй",
                "",
                "",
                "",
                "",
                "",
                "",
                "Түмт",
                "no.password",
                "",
                "",
            ],
            [
                "ТБ50000003",
                "Алдаа",
                "Гурав",
                "Эрэгтэй",
                "",
                "",
                "",
                "",
                "",
                "",
                "Түмт",
                "admin",
                "ImportPass2026",
                "",
            ],
        ],
        ACCOUNT_COLUMNS,
    )
    preview = _upload(client, csrf, content)
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["valid_rows"] == 0
    assert body["invalid_rows"] == 3
    fields = {error["field"] for error in body["errors"] if error["field"]}
    assert "password" in fields
    assert "username" in fields
    assert "short" not in preview.text
