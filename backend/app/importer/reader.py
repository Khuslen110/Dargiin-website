import csv
from collections.abc import Iterator
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from app.core.errors import bad_request


def _stringify(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value).strip()


def _csv_rows(path: Path) -> Iterator[tuple[int, list[str]]]:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            sample = handle.readline()
            handle.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(handle, dialect)
            for row_number, row in enumerate(reader, start=1):
                yield row_number, [_stringify(cell) for cell in row]
    except UnicodeDecodeError as exc:
        raise bad_request(
            "Файлын кодировк буруу байна. UTF-8 (BOM-той) CSV ашиглана уу."
        ) from exc
    except csv.Error as exc:
        raise bad_request("CSV файлыг уншиж чадсангүй. Бүтцийг шалгана уу.") from exc


def _xlsx_rows(path: Path) -> Iterator[tuple[int, list[str]]]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl-ийн олон төрлийн алдаа
        raise bad_request("Excel файлыг нээж чадсангүй. Файл гэмтсэн байж болно.") from exc
    try:
        sheet = workbook.active
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            yield row_number, [_stringify(cell) for cell in row]
    finally:
        workbook.close()


def iter_file_rows(path: Path) -> Iterator[tuple[int, list[str]]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        yield from _csv_rows(path)
    elif suffix in {".xlsx", ".xlsm"}:
        yield from _xlsx_rows(path)
    else:
        raise bad_request("Зөвхөн .csv эсвэл .xlsx файл оруулна уу.")


def read_header(path: Path) -> tuple[list[str], Iterator[tuple[int, list[str]]]]:
    rows = iter_file_rows(path)
    try:
        row_number, header = next(rows)
    except StopIteration as exc:
        raise bad_request("Файл хоосон байна.") from exc
    return header, rows
